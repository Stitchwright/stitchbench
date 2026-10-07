r"""Regenerate prisma/seed-data/service_modes.json from the owner's Service Modes sheet.

The sheet (C:\Manuals\_business-notes\Service_Modes.xlsx, first worksheet "Service modes") holds one row per procedure
family: how to put a machine into its test / service / diagnostic mode, how to leave it, and what the numbered tests
do, in our own words with the manual page. Edit the sheet, then run:

  python scripts/library/service_modes_from_xlsx.py            # sheet -> prisma/seed-data/service_modes.json
  python scripts/library/service_modes_from_xlsx.py --xlsx X.xlsx --out Y.json
  npm run db:seed                                              # loads the set into the ServiceMode table

Cell formats the script reads back:
  enter_steps / exit_steps   one step per line ("1. Hold Start/Stop and Reverse ..."); the numbers are optional
  tests                      "NN: purpose; NN: purpose" (a semicolon or a line break between tests; the first ": "
                             splits the code from the purpose, so keep semicolons out of a purpose)
  modelKeys                  leave blank to recompute from `models` (do that after editing `models`); otherwise the
                             comma-separated keys in the cell are used as they are
  pages                      1-based PDF page numbers ("196-197, 200"); the first one becomes `page`, where the
                             program opens the manual
  confidence                 high | medium | low
  figure_page, figure_box    optional columns for the picture of the entry keys the card shows beside the steps (cut
                             from the tech's own PDF in the browser, then shared like a factory setting's figure).
                             figure_page: a 1-based PDF page; blank = the first page in `pages`; "none" = no picture.
                             figure_box: "x0, y0, x1, y1" as fractions of the page (0 = left / top, 1 = right / bottom),
                             e.g. "0.08, 0.30, 0.55, 0.62"; blank = the whole page. Without the columns every row
                             shows its whole first page.
A row with an empty brand or no enter step is skipped. `id` is a short hash of brand, models, source file and mode,
so it changes when one of those is edited (the seed replaces the whole set, so that is harmless).

Rules for whoever edits the sheet: steps in our own words with the page reference, never copied manual text; never a
password, PIN, unlock code or lock-bypass step (write "needs the service password (owner keeps these privately)").
"""
import argparse, hashlib, json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
SEED = os.path.join(HERE, "..", "..", "prisma", "seed-data")
LIB = os.environ.get("MANUALS_ROOT", r"C:\Manuals")
DEFAULT_XLSX = os.path.join(LIB, "_business-notes", "Service_Modes.xlsx")
DEFAULT_OUT = os.path.join(SEED, "service_modes.json")

COLUMNS = ["brand", "models", "modelKeys", "category", "mode", "enter_steps", "exit_steps", "tests", "cautions",
           "source_file", "source_path", "pages", "confidence", "notes"]

# Words that are never a model on their own.
STOP = {"AND", "ALSO", "SERIES", "THE", "WITH", "SAME", "ONLY", "TYPE", "MODEL", "MODELS", "FACTORY", "CODES", "CODE",
        "PLUS", "NEEDLE", "NEEDLES", "STANDALONE", "CONTROL", "PANEL", "DISPLAY", "VERSION", "VERSIONS", "EARLIER",
        "COLOR", "COLOUR", "LCD", "TFT", "LED", "UNIT", "UNITS", "SUPPLEMENT", "SUPPLEMENTS", "LIKELY", "SIBLING",
        "EDITION", "PATCHWORK", "MACHINES", "SINGLE", "BRIDGE", "FAMILY", "HOBBYLOCK", "OPERATOR", "MANUAL", "MENU"}


def model_key(s: str) -> str:
    """Same rule as modelKey() in src/lib/format.ts."""
    return re.sub(r"[^A-Z0-9]", "", (s or "").upper())


def model_keys(models: str) -> list[str]:
    """Keys a machine's model can match: each listed model whole, each word with a digit in it (and its hyphen parts),
    and all-capital letter codes such as BLSO or BLG (Baby Lock models with no digit)."""
    keys = set()
    text = models or ""
    text = re.sub(r"\(factory codes?[^)]*\)", " ", text, flags=re.I)  # factory codes are not model names
    text = re.sub(r"\([^)0-9]*\)", " ", text)  # "(stand-alone type)", "(and Patchwork Edition)": no model in them

    def ok(k: str) -> bool:
        # a key needs a digit (bare numbers need three) or three letters, and is not a filler word
        if not 2 <= len(k) <= 24 or k in STOP:
            return False
        if k.isdigit():
            return len(k) >= 3
        return bool(re.search(r"\d", k)) or len(k) >= 3

    for part in re.split(r"[,;/+()]|\band\b|\s-\s", text):
        words = part.split()
        while words and model_key(words[0]) in STOP:  # "also 520" -> 520
            words.pop(0)
        while words and model_key(words[-1]) in STOP:  # "15 needle" -> 15
            words.pop()
        part = " ".join(words)
        if not part:
            continue
        whole = model_key(part)
        has_digit = bool(re.search(r"\d", part))
        if ok(whole) and (has_digit or (len(words) <= 3 and part != part.lower())):
            keys.add(whole)
        if len(words) >= 3 and has_digit:  # "Designer Jade 20" -> also JADE20
            tail = model_key(" ".join(words[-2:]))
            if ok(tail) and re.search(r"\d", tail):
                keys.add(tail)
        for w in words:
            if re.fullmatch(r"[A-Z]{3,}", w) and ok(w):  # Baby Lock codes with no digit: BLSO, BLG, BLL
                keys.add(w)
            if not re.search(r"\d", w):
                continue
            k = model_key(w)
            if ok(k):
                keys.add(k)
            for sub in w.split("-"):  # HZL-F300 -> F300, CS8060-CS8130 -> CS8060, CS8130
                ks = model_key(sub)
                if ok(ks) and re.search(r"\d", ks) and len(ks) >= 3:
                    keys.add(ks)
    return sorted(keys)


def row_id(brand: str, models: str, source_file: str, mode: str) -> str:
    return hashlib.sha1(f"{brand}|{models}|{source_file}|{mode}".encode("utf-8")).hexdigest()[:20]


def split_steps(cell) -> list[str]:
    text = str(cell or "").strip()
    if not text:
        return []
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    if len(lines) == 1:  # "1. a 2. b" typed on one line
        parts = re.split(r"(?:^|\s)\d{1,2}[.)]\s+", lines[0])
        lines = [p.strip() for p in parts if p.strip()] or lines
    return [re.sub(r"^\d{1,2}[.)]\s*", "", l).strip() for l in lines]


def split_tests(cell) -> list[dict]:
    out = []
    for item in re.split(r";|\n", str(cell or "")):
        item = item.strip()
        if not item:
            continue
        code, sep, purpose = item.partition(": ")
        if not sep:
            code, sep, purpose = item.partition(":")
        if not sep:
            code, purpose = "", item
        out.append({"code": code.strip(), "purpose": purpose.strip()})
    return out


def first_page(pages) -> int | None:
    m = re.search(r"\d+", str(pages or ""))
    return int(m.group(0)) if m else None


def figure_page(cell, page: int | None, where: str) -> int | None:
    """The picture's page: blank -> the procedure's first page, "none" (or "-", "0") -> no picture."""
    text = str(cell if cell is not None else "").strip()
    if not text:
        return page
    if text.lower() in ("none", "no", "-", "0"):
        return None
    m = re.fullmatch(r"(\d+)(?:\.0)?", text)  # a number cell comes back as 12 or 12.0
    if not m or int(m.group(1)) < 1:
        sys.exit(f"{where}: figure_page {text!r} is not a page number (or blank, or none)")
    return int(m.group(1))


def figure_box(cell, where: str) -> list[float]:
    """"x0, y0, x1, y1" as fractions of the page -> [x0, y0, x1, y1]; blank -> [] (the whole page)."""
    text = str(cell or "").strip()
    if not text:
        return []
    try:
        box = [float(x) for x in re.split(r"[,;\s]+", text.strip("[]() ")) if x]
    except ValueError:
        box = []
    if len(box) != 4 or not all(0 <= x <= 1 for x in box) or box[0] >= box[2] or box[1] >= box[3]:
        sys.exit(f"{where}: figure_box {text!r} must be four fractions x0, y0, x1, y1 with x0 < x1 and y0 < y1")
    return [round(x, 3) for x in box]  # the figure id rounds to three places too (figure_id in common.py)


def tests_cell(tests: list[dict]) -> str:
    """The sheet form of a test list (used by whoever writes the sheet from JSON)."""
    return ";\n".join(f"{t['code']}: {t['purpose'].replace(';', ',')}" if t.get("code") else t["purpose"].replace(";", ",")
                      for t in tests)


def steps_cell(steps: list[str]) -> str:
    return "\n".join(f"{i}. {s}" for i, s in enumerate(steps, 1))


def from_sheet(path: str) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Service modes"] if "Service modes" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(it)]
    idx = {h: i for i, h in enumerate(header)}
    missing = [c for c in COLUMNS if c not in idx]
    if missing:
        sys.exit(f"{path}: missing columns {missing}")
    rows, seen = [], {}
    for n, values in enumerate(it, start=2):
        # figure_page / figure_box are optional columns: a sheet without them reads as blank
        get = lambda c: values[idx[c]] if c in idx and idx[c] < len(values) else None
        brand = str(get("brand") or "").strip()
        enter = split_steps(get("enter_steps"))
        if not brand or not enter:
            continue
        models = str(get("models") or "").strip()
        source_file = str(get("source_file") or "").strip()
        mode = str(get("mode") or "").strip() or "Test mode"
        keys_cell = str(get("modelKeys") or "").strip()
        keys = sorted({model_key(k) for k in re.split(r"[,;\s]+", keys_cell) if model_key(k)}) if keys_cell else model_keys(models)
        rid = row_id(brand, models, source_file, mode)
        if rid in seen:
            sys.exit(f"{path}: rows {seen[rid]} and {n} have the same brand, models, source file and mode")
        seen[rid] = n
        pages = str(get("pages") or "").strip()
        page = first_page(pages)
        where = f"{path} row {n}"
        rows.append({
            "id": rid,
            "brand": brand,
            "models": models,
            "modelKeys": keys,
            "category": str(get("category") or "").strip() or None,
            "mode": mode,
            "enterSteps": enter,
            "exitSteps": split_steps(get("exit_steps")),
            "tests": split_tests(get("tests")),
            "cautions": str(get("cautions") or "").strip() or None,
            "sourceFile": source_file,
            "sourcePath": str(get("source_path") or "").strip() or None,
            "pages": pages or None,
            "page": page,
            "figurePage": figure_page(get("figure_page"), page, where),
            "figureBox": figure_box(get("figure_box"), where),
            "confidence": str(get("confidence") or "").strip().lower() or None,
            "notes": str(get("notes") or "").strip() or None,
        })
    return rows


def check(rows: list[dict]) -> None:
    """Warn about rows whose manual is not in the library or whose page is past the end of the PDF."""
    try:
        lib = {r["path"]: r for r in json.load(open(os.path.join(SEED, "library_docs.json"), encoding="utf-8"))}
    except FileNotFoundError:
        return
    for r in rows:
        p = r["sourcePath"]
        if not p:
            print(f"warning: {r['brand']} {r['models']}: no source_path", file=sys.stderr)
        elif p not in lib and not os.path.exists(os.path.join(LIB, p)):
            print(f"warning: {r['brand']} {r['models']}: {p} is not in the library", file=sys.stderr)
        elif p in lib and r["page"] and lib[p].get("pages") and r["page"] > lib[p]["pages"]:
            print(f"warning: {r['brand']} {r['models']}: page {r['page']} > {lib[p]['pages']} pages", file=sys.stderr)
        elif p in lib and r["figurePage"] and lib[p].get("pages") and r["figurePage"] > lib[p]["pages"]:
            print(f"warning: {r['brand']} {r['models']}: figure page {r['figurePage']} > {lib[p]['pages']} pages", file=sys.stderr)
        if not r["modelKeys"]:
            print(f"warning: {r['brand']} {r['models']}: no model keys", file=sys.stderr)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    ap.add_argument("--out", default=DEFAULT_OUT)
    a = ap.parse_args()
    rows = from_sheet(a.xlsx)
    check(rows)
    with open(a.out, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=1)
        f.write("\n")
    print(f"{len(rows)} service-mode rows -> {os.path.abspath(a.out)}")


if __name__ == "__main__":
    main()
