"""
Scan the manual library for distribution markings: "dealer only", "confidential", "internal use only", "not for
distribution" and the like. Read-only: it opens every PDF, reads the text of the first pages and the last page (where
covers, headers, footers and back pages carry such stamps) and writes a sheet for the owner. Scanned manuals with no
text layer can't be read this way; they are listed separately.

  python scripts/library/scan_markings.py [--root C:\\Manuals] [--pages 3] [--out <xlsx>]
"""
import argparse, csv, datetime, os, re, sys, time

import pymupdf as fitz  # PyMuPDF

STRONG = [
    ("confidential", r"\b(company\s+)?confidential\b"),
    ("dealer only", r"\b(for\s+)?(authori[sz]ed\s+)?dealers?(\s+use)?\s+only\b"),
    ("internal use only", r"\binternal\s+use\s+only\b|\bfor\s+internal\s+use\b"),
    ("not for distribution", r"\bnot\s+for\s+(public\s+)?(distribution|circulation|release|sale)\b|\bdo\s+not\s+(distribute|circulate)\b"),
    ("restricted", r"\brestricted\s+(document|information|distribution)\b"),
    ("proprietary", r"\bproprietary(\s+and\s+confidential|\s+information|\s+document)\b"),
]
MEDIUM = [
    ("service personnel only", r"\b(for\s+)?(authori[sz]ed\s+|qualified\s+)?(service|repair)\s+(personnel|technicians?|engineers?|centers?|centres?)\s+only\b"),
    ("not to the end user", r"\bnot\s+(to\s+be\s+)?(given|provided|supplied|distributed)\s+to\s+(the\s+)?(end\s+)?(users?|customers?|consumers?)\b"),
]
SKIP_DIRS = {"_business-notes", "_own-notes"}


def scan_pdf(path, pages):
    hits, text_found = [], False
    try:
        doc = fitz.open(path)
    except Exception as e:  # damaged or encrypted
        return None, f"could not open: {e.__class__.__name__}"
    try:
        if doc.needs_pass:
            return None, "password-protected"
        n = doc.page_count
        idx = sorted(set(list(range(min(pages, n))) + ([n - 1] if n else [])))
        for i in idx:
            try:
                t = doc.load_page(i).get_text("text") or ""
            except Exception:  # a damaged page (bad image data): skip it, keep the rest of the file
                continue
            if t.strip():
                text_found = True
            flat = re.sub(r"\s+", " ", t)
            for level, rules in (("strong", STRONG), ("medium", MEDIUM)):
                for label, rx in rules:
                    m = re.search(rx, flat, re.I)
                    if m:
                        a, b = max(0, m.start() - 50), min(len(flat), m.end() + 50)
                        hits.append((level, label, i + 1, flat[a:b].strip()))
        return (hits, text_found, n), None
    finally:
        doc.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=r"C:\Manuals")
    ap.add_argument("--pages", type=int, default=3)
    stamp = datetime.date.today().isoformat()
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    out = a.out or os.path.join(a.root, "_business-notes", f"Manual_Markings_{stamp}.xlsx")

    files = []
    for dirpath, dirnames, filenames in os.walk(a.root):
        dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS]
        for f in filenames:
            if f.lower().endswith(".pdf"):
                files.append(os.path.join(dirpath, f))
    files.sort()
    print(f"{len(files)} PDFs under {a.root}", flush=True)

    flagged, notext, problems = [], [], []
    t0 = time.time()
    for k, path in enumerate(files, 1):
        res, err = scan_pdf(path, a.pages)
        rel = os.path.relpath(path, a.root)
        parts = rel.split(os.sep)
        category = parts[0] if len(parts) > 1 else ""
        brand = parts[1] if len(parts) > 2 else ""
        if err:
            problems.append((rel, category, brand, err))
        else:
            hits, text_found, n = res
            if not text_found:
                notext.append((rel, category, brand, n))
            seen = set()
            for level, label, page, snip in hits:
                if (label, page) in seen:
                    continue
                seen.add((label, page))
                flagged.append((rel, category, brand, level, label, page, snip[:160]))
        if k % 500 == 0:
            print(f"  {k}/{len(files)} in {time.time() - t0:.0f}s, {len({f[0] for f in flagged})} files flagged", flush=True)

    files_flagged = {}
    for row in flagged:
        files_flagged.setdefault(row[0], []).append(row)
    strong_files = {f for f, rows in files_flagged.items() if any(r[3] == "strong" for r in rows)}

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "Flagged"
        ws.append(["file", "category", "brand", "level", "marking", "page", "context"])
        for r in sorted(flagged, key=lambda r: (r[3] != "strong", r[2].lower(), r[0], r[5])):
            ws.append(list(r))
        ws2 = wb.create_sheet("By brand")
        ws2.append(["brand", "files with a strong marking", "files with a medium marking only", "PDFs scanned for that brand"])
        by_brand = {}
        for path in files:
            parts = os.path.relpath(path, a.root).split(os.sep)
            b = parts[1] if len(parts) > 2 else "(top level)"
            by_brand.setdefault(b, [0, 0, 0])[2] += 1
        for f, rows in files_flagged.items():
            parts = f.split(os.sep)
            b = parts[1] if len(parts) > 2 else "(top level)"
            if f in strong_files:
                by_brand[b][0] += 1
            else:
                by_brand[b][1] += 1
        for b, (s, m, n) in sorted(by_brand.items(), key=lambda kv: (-kv[1][0], -kv[1][1], kv[0].lower())):
            ws2.append([b, s, m, n])
        ws3 = wb.create_sheet("No text (scans)")
        ws3.append(["file", "category", "brand", "pages"])
        for r in notext:
            ws3.append(list(r))
        ws4 = wb.create_sheet("Could not open")
        ws4.append(["file", "category", "brand", "why"])
        for r in problems:
            ws4.append(list(r))
        ws5 = wb.create_sheet("About")
        for line in [
            f"Scanned {stamp}: {len(files)} PDFs under {a.root} (the business-notes and own-notes folders skipped).",
            f"Pages read per PDF: the first {a.pages} and the last. Markings in the middle of a manual only are not seen.",
            "Strong: confidential, dealer only, internal use only, not for distribution, restricted, proprietary and confidential.",
            "Medium: for service personnel / technicians / service centers only; not to be given to end users.",
            "A copyright line ('All rights reserved', 'may not be reproduced') is NOT flagged: nearly every manual has one.",
            "No text (scans): image-only PDFs; their markings can't be read without OCR.",
            "Read-only: no file in the library was changed.",
        ]:
            ws5.append([line])
        for w in wb.worksheets:
            w["A1"].font = Font(bold=True)
        wb.save(out)
    except ImportError:
        out = out.replace(".xlsx", ".csv")
        with open(out, "w", newline="", encoding="utf-8") as fh:
            csv.writer(fh).writerows([["file", "category", "brand", "level", "marking", "page", "context"], *flagged])

    print(f"done in {time.time() - t0:.0f}s")
    print(f"PDFs scanned: {len(files)}")
    print(f"files with a strong marking: {len(strong_files)}")
    print(f"files with a medium marking only: {len(files_flagged) - len(strong_files)}")
    print(f"image-only (no text): {len(notext)}")
    print(f"could not open: {len(problems)}")
    print(f"sheet: {out}")


if __name__ == "__main__":
    sys.exit(main())
